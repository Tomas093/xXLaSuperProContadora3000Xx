from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS_SKILL = Path("/Users/pedrodelaguila/.codex/plugins/cache/openai-primary-runtime/documents/26.430.10722/skills/documents")
sys.path.append(str(DOCS_SKILL / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402

OUTPUT_BASE = ROOT / "proyecto/test_data/outputs/investigacion-2"
CATALOG_ANALYSIS = ROOT / "proyecto/symbol-catalog/investigacion_2/dorrego/catalog_visual_analysis.json"
CATALOG_OUTPUT_BASE = OUTPUT_BASE / "dorrego/catalog-visual-analysis/v1"
DEST = ROOT / "proyecto/test_data/reports/Investigación Parte II Deteccion con LLM.docx"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def run_key(path: Path) -> str:
    rel = path.relative_to(OUTPUT_BASE)
    return str(rel.parent)


def display_run(key: str) -> str:
    cleaned = key.replace("dorrego/", "")
    cleaned = cleaned.replace("claude-sonnet", "Sonnet")
    cleaned = cleaned.replace("claude-opus", "Opus")
    cleaned = cleaned.replace("entrepiso/", "")
    cleaned = cleaned.replace("firsphotos", "first photos")
    cleaned = cleaned.replace("firstphotos", "first photos")
    cleaned = cleaned.replace("second-photos", "second photos")
    return cleaned


def normalize_text(value) -> str:
    if value is None:
        return "-"
    text = str(value).strip()
    return text if text else "-"


def read_text(path: Path | None) -> str:
    if not path or not path.exists():
        return "-"
    return path.read_text(encoding="utf-8").strip() or "-"


def collect_runs() -> list[dict]:
    runs: dict[str, dict] = defaultdict(dict)
    for path in sorted(OUTPUT_BASE.rglob("*_bom.json")):
        runs[run_key(path)]["bom_path"] = path
        runs[run_key(path)]["bom"] = read_json(path)
    for path in sorted(OUTPUT_BASE.rglob("*_usage.json")):
        if "catalog-visual-analysis" in str(path):
            continue
        runs[run_key(path)]["usage_path"] = path
        runs[run_key(path)]["usage"] = read_json(path)
    for path in sorted(OUTPUT_BASE.rglob("*_raw.txt")):
        if "catalog-visual-analysis" in str(path):
            continue
        runs[run_key(path)]["raw_path"] = path
    for path in sorted(OUTPUT_BASE.rglob("*_bom.csv")):
        runs[run_key(path)]["csv_path"] = path

    records = []
    for key, data in sorted(runs.items()):
        if "bom" not in data:
            continue
        bom = data["bom"]
        usage = data.get("usage", {})
        total_usage = usage.get("usage", {}).get("total", {})
        cost = usage.get("estimated_cost_usd", {}).get("total", {})
        records.append(
            {
                "key": key,
                "display": display_run(key),
                "model": usage.get("model", "-"),
                "provider": usage.get("provider", "-"),
                "bom_rows": bom.get("bom", []),
                "unknowns": bom.get("simbolos_no_identificados", []),
                "input_tokens": total_usage.get("input_tokens", 0),
                "output_tokens": total_usage.get("output_tokens", 0),
                "cost": cost.get("total"),
                "paths": {
                    "bom": data.get("bom_path"),
                    "csv": data.get("csv_path"),
                    "raw": data.get("raw_path"),
                    "usage": data.get("usage_path"),
                },
            }
        )
    return records


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_code_block(doc: Document, text: str, chunk_size: int = 1800) -> None:
    content = normalize_text(text)
    for start in range(0, len(content), chunk_size):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(content[start : start + chunk_size])
        run.font.name = "Courier New"
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(40, 40, 40)


def style_table(table, widths: list[int], header_fill: str = "E8EEF7", font_size: int = 8) -> None:
    table.style = "Table Grid"
    apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=0)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], font_size: int = 8) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, size=font_size)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            set_cell_text(cells[idx], normalize_text(value), size=font_size)
    style_table(table, widths, font_size=font_size)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for style_name, size, color in [
        ("Title", 22, RGBColor(31, 78, 121)),
        ("Heading 1", 15, RGBColor(31, 78, 121)),
        ("Heading 2", 12, RGBColor(47, 47, 47)),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "LIARD - Investigación Parte II"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.text = "Deteccion con LLM"
    for run in fp.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)


def add_intro(doc: Document, runs: list[dict]) -> None:
    doc.add_heading("Investigación Parte II Deteccion con LLM", 0)
    subtitle = doc.add_paragraph()
    subtitle.add_run("Compilacion de resultados de investigacion 2 - Dorrego / Entrepiso").bold = True
    doc.add_paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(
        "Este documento consolida los outputs disponibles de investigacion 2: BOM normalizados, simbolos no identificados, costos/uso por corrida, y el analisis visual del catalogo usado para mejorar la diferenciacion de simbolos."
    )

    total_cost = sum(float(r["cost"] or 0) for r in runs)
    doc.add_heading("Resumen ejecutivo", 1)
    add_table(
        doc,
        ["Metrica", "Valor"],
        [
            ["Corridas BOM analizadas", str(len(runs))],
            ["Costo total de corridas BOM registradas", f"USD {total_cost:.6f}"],
            ["Catalog visual analysis disponible", "Si" if CATALOG_ANALYSIS.exists() else "No"],
            ["Catalog visual analysis raw", "Si" if (CATALOG_OUTPUT_BASE / "raw.txt").exists() else "No"],
            ["Carpeta base de outputs", str(OUTPUT_BASE.relative_to(ROOT))],
        ],
        [3600, 5760],
        font_size=9,
    )


def add_runs_summary(doc: Document, runs: list[dict]) -> None:
    doc.add_heading("Resumen de corridas", 1)
    rows = []
    for run in runs:
        rows.append(
            [
                run["display"],
                run["model"],
                str(len(run["bom_rows"])),
                str(len(run["unknowns"])),
                str(run["input_tokens"]),
                str(run["output_tokens"]),
                f"{float(run['cost'] or 0):.6f}",
            ]
        )
    add_table(
        doc,
        ["Corrida", "Modelo", "BOM", "No ident.", "Input tok.", "Output tok.", "USD"],
        rows,
        [2500, 1700, 650, 800, 1050, 1050, 800],
        font_size=7,
    )


def add_usage_details(doc: Document, runs: list[dict]) -> None:
    doc.add_heading("Uso, costos y prompts por corrida", 1)
    rows = []
    for run in runs:
        usage = read_json(run["paths"]["usage"]) if run["paths"].get("usage") else {}
        prompts = usage.get("prompts", {})
        bom_usage = usage.get("usage", {}).get("bom_generation", {})
        timings = usage.get("timings_seconds", {})
        pricing = usage.get("pricing_usd_per_million_tokens", {})
        rows.append(
            [
                run["display"],
                usage.get("provider", "-"),
                usage.get("model", "-"),
                bom_usage.get("input_tokens", "-"),
                bom_usage.get("output_tokens", "-"),
                timings.get("total", "-"),
                f"{float(run['cost'] or 0):.6f}",
                prompts.get("plan_system", {}).get("path", "-"),
                prompts.get("plan_user", {}).get("path", "-"),
                pricing.get("input", "-"),
                pricing.get("output", "-"),
            ]
        )
    add_table(
        doc,
        [
            "Corrida",
            "Provider",
            "Modelo",
            "Input",
            "Output",
            "Seg.",
            "USD",
            "Plan system",
            "Plan user",
            "USD/M input",
            "USD/M output",
        ],
        rows,
        [1500, 700, 1200, 650, 650, 550, 650, 1650, 1650, 650, 650],
        font_size=6,
    )


def add_bom_details(doc: Document, runs: list[dict]) -> None:
    doc.add_heading("Detalle BOM por corrida", 1)
    rows = []
    for run in runs:
        for item in run["bom_rows"]:
            rows.append(
                [
                    run["display"],
                    item.get("material", ""),
                    item.get("especificacion"),
                    str(item.get("cantidad", "")),
                ]
            )
    add_table(
        doc,
        ["Corrida", "Material", "Especificacion", "Cantidad"],
        rows,
        [2500, 2600, 3460, 800],
        font_size=7,
    )


def add_unknowns(doc: Document, runs: list[dict]) -> None:
    doc.add_heading("Simbolos no identificados", 1)
    rows = []
    for run in runs:
        for item in run["unknowns"]:
            rows.append(
                [
                    run["display"],
                    item.get("material_mas_parecido"),
                    item.get("ubicacion_aproximada"),
                    item.get("texto_derecha_posible_especificacion"),
                    item.get("motivo"),
                ]
            )
    if not rows:
        doc.add_paragraph("No hay simbolos no identificados en los outputs analizados.")
        return
    add_table(
        doc,
        ["Corrida", "Mas parecido", "Ubicacion", "Texto derecha", "Motivo"],
        rows,
        [1900, 1700, 2300, 1500, 1960],
        font_size=6,
    )


def add_catalog_analysis(doc: Document) -> None:
    doc.add_heading("Catalog visual analysis", 1)
    if not CATALOG_ANALYSIS.exists():
        doc.add_paragraph("No se encontro catalog_visual_analysis.json.")
        return

    data = read_json(CATALOG_ANALYSIS)
    rows = []
    for item in data.get("catalog_visual_analysis", []):
        rows.append(
            [
                item.get("component_name", ""),
                item.get("filename", ""),
                item.get("visual_signature", ""),
                "; ".join(item.get("distinguishing_features", [])),
                "; ".join(item.get("negative_identification_rules", [])),
            ]
        )
    add_table(
        doc,
        ["Componente", "Archivo", "Firma visual", "Rasgos distintivos", "Reglas negativas"],
        rows,
        [1600, 1500, 2460, 2200, 1600],
        font_size=6,
    )

    raw_path = CATALOG_OUTPUT_BASE / "raw.txt"
    usage_path = CATALOG_OUTPUT_BASE / "usage.json"
    doc.add_heading("Catalog visual analysis - usage", 2)
    if usage_path.exists():
        usage = read_json(usage_path)
        total = usage.get("usage", {}).get("total", {})
        cost = usage.get("estimated_cost_usd", {}).get("total", {})
        add_table(
            doc,
            ["Modelo", "Input tok.", "Output tok.", "Costo USD", "Ruta"],
            [
                [
                    usage.get("model", "-"),
                    total.get("input_tokens", "-"),
                    total.get("output_tokens", "-"),
                    f"{float(cost.get('total') or 0):.6f}",
                    str(usage_path.relative_to(ROOT)),
                ]
            ],
            [1800, 1200, 1200, 1200, 3960],
            font_size=7,
        )
    else:
        doc.add_paragraph("No se encontro usage.json de catalog visual analysis.")

    doc.add_heading("Catalog visual analysis - raw output", 2)
    add_code_block(doc, read_text(raw_path))


def add_raw_outputs(doc: Document, runs: list[dict]) -> None:
    doc.add_heading("Raw outputs completos por corrida", 1)
    for run in runs:
        doc.add_heading(run["display"], 2)
        doc.add_paragraph(str(run["paths"].get("raw", "-")).replace(str(ROOT) + "/", ""))
        add_code_block(doc, read_text(run["paths"].get("raw")))


def add_artifact_paths(doc: Document, runs: list[dict]) -> None:
    doc.add_heading("Archivos fuente", 1)
    rows = []
    for run in runs:
        for kind, path in run["paths"].items():
            if path:
                rows.append([run["display"], kind, str(path.relative_to(ROOT))])
    rows.append(["Catalog visual analysis", "json", str(CATALOG_ANALYSIS.relative_to(ROOT))])
    raw_path = CATALOG_OUTPUT_BASE / "raw.txt"
    usage_path = CATALOG_OUTPUT_BASE / "usage.json"
    if raw_path.exists():
        rows.append(["Catalog visual analysis", "raw", str(raw_path.relative_to(ROOT))])
    if usage_path.exists():
        rows.append(["Catalog visual analysis", "usage", str(usage_path.relative_to(ROOT))])
    add_table(doc, ["Corrida", "Tipo", "Ruta"], rows, [2600, 900, 5860], font_size=7)


def main() -> None:
    runs = collect_runs()
    DEST.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_document(doc)
    add_intro(doc, runs)
    add_runs_summary(doc, runs)
    doc.add_page_break()
    add_usage_details(doc, runs)
    doc.add_page_break()
    add_bom_details(doc, runs)
    doc.add_page_break()
    add_unknowns(doc, runs)
    doc.add_page_break()
    add_catalog_analysis(doc)
    doc.add_page_break()
    add_raw_outputs(doc, runs)
    doc.add_page_break()
    add_artifact_paths(doc, runs)
    doc.save(DEST)
    print(DEST)


if __name__ == "__main__":
    main()
